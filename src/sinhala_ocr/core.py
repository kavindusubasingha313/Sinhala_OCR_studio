"""Document storage and engine contracts, independent of the interface."""
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Protocol
import json
import os
import tempfile

from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from .equations import Equation


@dataclass
class Recognition:
    text: str
    engine: str
    # Unknown confidence is distinct from high confidence.
    confidence: float | None = None


class OcrEngine(Protocol):
    def recognize(self, image: Path) -> Recognition: ...


class TesseractEngine:
    def recognize(self, image: Path) -> Recognition:
        import pytesseract
        available = set(pytesseract.get_languages())
        if not {"sin", "eng"} <= available:
            raise RuntimeError("Install Tesseract Sinhala (sin) and English (eng) language data.")
        with Image.open(image) as source:
            return Recognition(pytesseract.image_to_string(source, lang="sin+eng"), "tesseract")


@dataclass
class DocumentPage:
    source_index: int = 0
    text: str = ""
    approved: bool = False
    equations: list[Equation] = field(default_factory=list)
    history: list[dict] = field(default_factory=list)


@dataclass
class Document:
    source: str
    text: str = ""
    approved: bool = False
    latex_enabled: bool = False
    equations: list[Equation] = field(default_factory=list)
    pages: list[DocumentPage] = field(default_factory=list)
    active_page: int = 0

    def capture_page(self):
        if self.pages:
            page = self.pages[self.active_page]
            page.text, page.approved, page.equations = self.text, self.approved, self.equations

    def activate_page(self, index):
        if not 0 <= index < len(self.pages):
            raise ValueError("Page does not exist.")
        if index != self.active_page:
            self.capture_page()
        self.active_page = index
        page = self.pages[index]
        self.text, self.approved, self.equations = page.text, page.approved, page.equations

    def for_export(self):
        from copy import deepcopy
        self.capture_page()
        if not self.pages:
            return deepcopy(self)
        return Document(self.source, "\n\n".join(page.text for page in self.pages),
                        all(page.approved for page in self.pages), self.latex_enabled,
                        deepcopy([eq for page in self.pages for eq in page.equations]))

    def apply_recognition(self, result: Recognition) -> None:
        # Every edit is protected, even before explicit approval.
        if self.text or self.approved:
            raise ValueError("Existing text is protected. Open a new document to rerun OCR.")
        self.text = result.text


def import_document(path: Path) -> Document:
    from .documents import import_pages
    return import_pages(path)


def atomic_write(destination: Path, writer, validator) -> None:
    handle, temporary = tempfile.mkstemp(dir=destination.parent, suffix=destination.suffix)
    os.close(handle)
    path = Path(temporary)
    try:
        writer(path)
        validator(path)
        # Windows FlushFileBuffers requires a handle with write access.
        # r+b retains the validated bytes without truncating the temporary file.
        with path.open("r+b") as stream:
            os.fsync(stream.fileno())
        os.replace(path, destination)
    finally:
        path.unlink(missing_ok=True)


def save_project(document: Document, destination: Path) -> None:
    if Path(document.source).resolve() == destination.resolve():
        raise ValueError("Project cannot overwrite the original source.")
    document.capture_page()
    atomic_write(destination,
                 lambda p: p.write_text(json.dumps(asdict(document), ensure_ascii=False), encoding="utf-8"),
                 lambda p: load_project(p))


def load_project(path: Path) -> Document:
    data = json.loads(path.read_text(encoding="utf-8"))
    data["equations"] = [Equation(**equation) for equation in data.get("equations", [])]
    data["pages"] = [DocumentPage(**{**page, "equations": [Equation(**eq) for eq in page.get("equations", [])]}) for page in data.get("pages", [])]
    document = Document(**data)
    if document.pages:
        document.activate_page(document.active_page)
    return document


def export_document(document: Document, destination: Path, *, image_fallback=False) -> None:
    document = document.for_export()
    if Path(document.source).resolve() == destination.resolve():
        raise ValueError("Export cannot overwrite the source document.")
    if document.equations or "[[eq:" in document.text or destination.suffix.lower() in {".md", ".tex"}:
        from .equation_export import export_with_equations
        export_with_equations(document, destination, image_fallback=image_fallback)
        return
    if destination.suffix.lower() == ".txt":
        def validate(path):
            if path.read_text(encoding="utf-8") != document.text:
                raise ValueError("Text integrity check failed")
        atomic_write(destination, lambda p: p.write_text(document.text, encoding="utf-8"), validate)
    elif destination.suffix.lower() == ".docx":
        def write(path):
            word = WordDocument()
            for line in document.text.split("\n"):
                run = word.add_paragraph().add_run(line)
                run.font.name = "Nirmala UI"
                fonts = run._element.get_or_add_rPr().rFonts
                fonts.set(qn("w:cs"), "Nirmala UI")
                language = OxmlElement("w:lang")
                language.set(qn("w:bidi"), "si-LK")
                run._element.get_or_add_rPr().append(language)
            word.save(path)
        def validate(path):
            word = WordDocument(path)
            if "\n".join(p.text for p in word.paragraphs) != document.text:
                raise ValueError("DOCX text integrity check failed")
        atomic_write(destination, write, validate)
    else:
        raise ValueError("Stage 1 supports TXT and DOCX exports only.")
