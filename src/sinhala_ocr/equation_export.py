"""Explicit equation placement and validated exports; no silent flattening."""
from io import BytesIO
import base64
import re
from zipfile import ZipFile

from docx import Document as WordDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree

from .equations import EquationError, equation_omml

MARKER = re.compile(r"\[\[eq:([a-f0-9]{32})\]\]")


def document_parts(document):
    equations = {equation.id: equation for equation in document.equations}
    if len(equations) != len(document.equations):
        raise EquationError("Duplicate equation IDs in project.")
    parts, cursor, used = [], 0, set()
    for match in MARKER.finditer(document.text):
        parts.append(document.text[cursor:match.start()])
        identity = match.group(1)
        if identity not in equations or identity in used:
            raise EquationError("Equation marker is missing its data or occurs twice. Undo the text edit or restore its marker.")
        equation = equations[identity]
        if equation.status == "rejected":
            raise EquationError("A rejected equation still has a marker. Restore it or remove its marker before export.")
        parts.append(equation)
        used.add(identity)
        cursor = match.end()
    parts.append(document.text[cursor:])
    if any("[[eq:" in part for part in parts if isinstance(part, str)):
        raise EquationError("An equation marker is incomplete. Undo or correct the marker before export.")
    if any(eq.status != "rejected" and eq.id not in used for eq in document.equations):
        raise EquationError("An equation is missing from the text. Reinsert its marker or reject it before export.")
    return parts


def editable_equation_errors(document):
    errors = []
    for part in document_parts(document):
        if not isinstance(part, str):
            try:
                equation_omml(part)
            except EquationError as error:
                errors.append((part, str(error)))
            else:
                if part.status != "accepted":
                    raise EquationError("Review and Accept this equation before exporting.")
    return errors


def export_with_equations(document, destination, *, image_fallback=False):
    from .core import atomic_write
    parts = document_parts(document)
    extension = destination.suffix.lower()
    if extension in {".md", ".tex"}:
        # Source export keeps even unsupported LaTeX exactly; it does not claim to compile it.
        def escape_text(text):
            if extension == ".md":
                return text.replace("\\", "\\\\").replace("$", "\\$")
            return "".join({"\\": r"\textbackslash{}", "{": r"\{", "}": r"\}",
                            "$": r"\$", "&": r"\&", "#": r"\#", "%": r"\%",
                            "_": r"\_", "^": r"\textasciicircum{}", "~": r"\textasciitilde{}"}.get(c, c) for c in text)
        output = []
        for part in parts:
            if isinstance(part, str):
                output.append(escape_text(part))
            elif extension == ".md":
                output.append(f"${part.latex}$" if part.placement == "inline" else f"\n\n$$\n{part.latex}\n$$\n\n")
            else:
                output.append(r"\(" + part.latex + r"\)" if part.placement == "inline" else "\n\\[" + part.latex + "\\]\n")
        content = "".join(output)
        def validate(path):
            if path.read_text(encoding="utf-8") != content:
                raise EquationError("Equation source export integrity check failed.")
        atomic_write(destination, lambda p: p.write_text(content, encoding="utf-8"), validate)
        return
    if extension != ".docx":
        raise EquationError("Documents with equations support DOCX, Markdown and LaTeX source export. TXT/PDF equation export is not available; no equations were discarded.")
    converted, fallbacks = {}, {}
    for part in parts:
        if isinstance(part, str):
            continue
        try:
            converted[part.id] = equation_omml(part)
        except EquationError:
            if not image_fallback or not part.image_png:
                raise EquationError("This equation cannot become an editable Word object. Correct it, export its LaTeX source, or explicitly choose an original-image fallback when a source crop exists.")
            fallbacks[part.id] = base64.b64decode(part.image_png, validate=True)
        else:
            if part.status != "accepted":
                raise EquationError("Review and Accept every supported equation before DOCX export.")
    label = "[Equation image fallback — not editable]"
    # Compare content in actual order, including explicit fallback labels.
    expected_text = "".join(part.replace("\n", "") if isinstance(part, str) else (label if part.id in fallbacks else "") for part in parts)

    def write(path):
        word = WordDocument()
        math_properties = word.settings.element.find(qn("m:mathPr"))
        if math_properties is None:
            math_properties = OxmlElement("m:mathPr")
            word.settings.element.insert_element_before(math_properties, "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping", "w:decimalSymbol", "w:listSeparator")
        math_font = math_properties.find(qn("m:mathFont"))
        if math_font is None:
            math_font = OxmlElement("m:mathFont")
            math_properties.insert(0, math_font)
        math_font.set(qn("m:val"), "Cambria Math")
        paragraph = word.add_paragraph()
        for part in parts:
            if isinstance(part, str):
                for index, line in enumerate(part.split("\n")):
                    if index:
                        paragraph = word.add_paragraph()
                    run = paragraph.add_run(line)
                    run.font.name = "Nirmala UI"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Nirmala UI")
            else:
                if part.placement == "display" and len(paragraph._p):
                    paragraph = word.add_paragraph()
                if part.id in fallbacks:
                    paragraph.add_run().add_picture(BytesIO(fallbacks[part.id]), width=Inches(3))
                    paragraph.add_run(label)
                elif part.placement == "display":
                    block = OxmlElement("m:oMathPara")
                    block.append(converted[part.id])
                    paragraph._p.append(block)
                else:
                    paragraph._p.append(converted[part.id])
                if part.placement == "display":
                    paragraph = word.add_paragraph()
        word.save(path)

    def validate(path):
        with ZipFile(path) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
        if "".join(root.xpath("//w:t/text()", namespaces=ns)) != expected_text:
            raise EquationError("DOCX surrounding text integrity failed.")
        if len(root.xpath("//m:oMath", namespaces=ns)) != len(converted):
            raise EquationError("DOCX editable equation count failed.")
        if root.xpath("//w:txbxContent", namespaces=ns):
            raise EquationError("Unexpected text boxes in equation export.")
        display_count = sum(not isinstance(part, str) and part.placement == "display" and part.id not in fallbacks for part in parts)
        if len(root.xpath("//m:oMathPara", namespaces=ns)) != display_count:
            raise EquationError("DOCX equation placement validation failed.")
        if len(root.xpath("//w:drawing", namespaces=ns)) != len(fallbacks):
            raise EquationError("DOCX image fallback validation failed.")
        WordDocument(path)  # Independently reopen the archive.
    atomic_write(destination, write, validate)
