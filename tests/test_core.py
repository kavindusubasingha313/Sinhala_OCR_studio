from zipfile import ZipFile

import pytest
from PIL import Image
from docx import Document as WordDocument

from sinhala_ocr.core import (Document, Recognition, atomic_write, export_document,
                             import_document, load_project, save_project)

TEXT = "සිංහල ශ්‍රී ලංකාව ක්‍රියා කු කූ කෙ කේ කෛ කො කෝ කෞ\nEnglish 123 + x² = ½"


@pytest.mark.parametrize("extension", [".txt", ".docx"])
def test_export_unicode(tmp_path, extension):
    path = tmp_path / ("output" + extension)
    export_document(Document("original.png", TEXT), path)
    if extension == ".txt":
        assert path.read_text(encoding="utf-8") == TEXT
    else:
        assert "\n".join(p.text for p in WordDocument(path).paragraphs) == TEXT
        with ZipFile(path) as archive:
            xml = archive.read("word/document.xml")
        assert b"txbxContent" not in xml
        assert b'w:cs="Nirmala UI"' in xml


def test_project_and_correction_protection(tmp_path):
    document = Document("image.png", TEXT, True)
    path = tmp_path / "project.json"
    save_project(document, path)
    restored = load_project(path)
    assert restored == document
    with pytest.raises(ValueError):
        restored.apply_recognition(Recognition("replacement", "fake"))
    assert restored.text == TEXT


def test_failed_validation_preserves_destination(tmp_path):
    path = tmp_path / "text.txt"
    path.write_text("existing")
    def fail(_):
        raise ValueError("invalid")
    with pytest.raises(ValueError):
        atomic_write(path, lambda p: p.write_text("broken"), fail)
    assert path.read_text() == "existing"
    assert list(tmp_path.iterdir()) == [path]


def test_source_cannot_be_overwritten(tmp_path):
    path = tmp_path / "original.txt"
    path.write_text("original")
    with pytest.raises(ValueError):
        export_document(Document(str(path), "replacement"), path)
    assert path.read_text() == "original"


def test_image_import(tmp_path):
    path = tmp_path / "image.png"
    Image.new("RGB", (30, 30), "white").save(path)
    assert import_document(path) == Document(str(path))


def test_reject_unsupported_export(tmp_path):
    with pytest.raises(ValueError):
        export_document(Document("source", TEXT), tmp_path / "out.pdf")


def test_selectable_pdf_import(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import (DecodedStreamObject, DictionaryObject, NameObject)
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 200 Td (Selectable English 123) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    path = tmp_path / "selectable.pdf"
    writer.write(path)
    assert "Selectable English 123" in import_document(path).text


def test_real_tesseract_adapter(tmp_path):
    from sinhala_ocr.core import TesseractEngine
    path = tmp_path / "blank.png"
    Image.new("RGB", (500, 200), "white").save(path)
    result = TesseractEngine().recognize(path)
    assert result.engine == "tesseract"
    assert result.text.strip() == ""
    assert result.confidence is None


def test_real_english_ocr_regression(tmp_path):
    from PIL import ImageDraw, ImageFont
    from sinhala_ocr.core import TesseractEngine
    path = tmp_path / "english.png"
    image = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(image).text((30, 40), "TEST 123", font=ImageFont.load_default(size=72), fill="black")
    image.save(path)
    assert TesseractEngine().recognize(path).text.strip() == "TEST 123"
